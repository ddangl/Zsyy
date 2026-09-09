import pandas as pd

import openpyxl

import os

import re

import sys

from datetime import datetime, timedelta

from jinja2 import Template

from tabulate import tabulate



try:

    from playwright.sync_api import sync_playwright

    PLAYWRIGHT_AVAILABLE = True

except ImportError:

    PLAYWRIGHT_AVAILABLE = False



import config





def convert_to_half_width(text):

    if pd.isna(text):

        return text

    text = str(text)

    text = text.translate(str.maketrans('０１２３４５６７８９：', '0123456789:'))

    text = text.strip()

    return text





def normalize_time_str(text, is_start_time=False):

    if pd.isna(text):

        return text

    text = str(text)
    text = re.sub(r'[.。·]', ':', text)
    text = text.replace('：', ':')
    text = re.sub(r'点', ':', text)
    text = re.sub(r'分', '', text)
    text = re.sub(r'\s*:\s*', ':', text)
    text = text.strip()



    match = re.match(r'^(\d{1,2}):(\d{1,2})(?::\d{1,2})?$', text)

    if match:

        hour = int(match.group(1))

        minute = int(match.group(2))

        if not is_start_time and hour < 10:

            hour = hour + 24

        return f"{hour:02d}:{minute:02d}"



    match_hour_only = re.match(r'^(\d{1,2}):$', text)

    if match_hour_only:

        hour = int(match_hour_only.group(1))

        if not is_start_time and hour < 10:

            hour = hour + 24

        return f"{hour:02d}:00"



    return text





def _check_time_valid(hours_str, minutes_str):

    try:

        hours = int(hours_str)

        minutes = int(minutes_str)

        return hours >= 0 and 0 <= minutes <= 59

    except:

        return False





def is_valid_time(text):

    if pd.isna(text):

        return False

    text = str(text)



    if '上班' in text or '开始' in text:

        return True



    if '(' in text or '（' in text:

        parsed = parse_time_string(text)

        if parsed:

            main_time = parsed.get('main_time', '')

            start_time = parsed.get('start_time')



            if main_time and main_time.strip():

                normalized_main = normalize_time_str(main_time)

                if normalized_main and ':' in normalized_main:

                    parts = normalized_main.split(':')

                    if _check_time_valid(parts[0], parts[1]):

                        return True



            if start_time:

                parts = start_time.split(':')

                if _check_time_valid(parts[0], parts[1]):

                    return True

        return False



    normalized_time = normalize_time_str(text)

    if normalized_time and ':' in normalized_time:

        parts = normalized_time.split(':')

        if _check_time_valid(parts[0], parts[1]):

            return True

    return False





def parse_time_string(time_str):

    if not time_str or pd.isna(time_str):

        return None



    time_str = str(time_str).strip()



    main_time = time_str.split('（')[0].split('(')[0].strip()



    start_time = None

    start_time_match = re.search(r'（(\d{1,2})[:：点](\d{1,2})分?(?:开始|上班|到)?）', time_str)

    if not start_time_match:

        start_time_match = re.search(r'\((\d{1,2})[:：点](\d{1,2})分?(?:开始|上班|到)?\)', time_str)

    if not start_time_match:

        start_time_match = re.search(r'[（(](\d{1,2})点(\d{1,2})分?(?:开始|上班|到)?[）)]', time_str)

    if not start_time_match:

        start_time_match = re.search(r'[（(](\d{1,2})点(?:开始|上班|到)?[）)]', time_str)



    if start_time_match:

        if start_time_match.lastindex == 2:

            start_time = f"{int(start_time_match.group(1)):02d}:{int(start_time_match.group(2)):02d}"

        else:

            start_time = f"{int(start_time_match.group(1)):02d}:00"



    main_time = convert_to_half_width(main_time)



    return {

        'main_time': main_time,

        'start_time': start_time

    }





def get_time_category(time_str, work_hours):

    if isinstance(work_hours, (int, float)):

        for category, (low, high) in config.WORK_HOURS_THRESHOLDS.items():

            if low <= work_hours < high:

                return category

    return 'normal'





def _find_target_sheets(file_path):

    target_col = config.TIME_OFF_WORK_COL

    xl = pd.ExcelFile(file_path)

    result = []

    for name in xl.sheet_names:

        df = pd.read_excel(file_path, sheet_name=name, nrows=1)

        if target_col in df.columns:

            result.append(name)

    return result





def _filter_by_checkin_time(df, before_hour=9):
    if '打卡时间' not in df.columns:
        return df, pd.DataFrame(columns=df.columns)
    df_copy = df.copy()
    df_copy['_dt'] = pd.to_datetime(df_copy['打卡时间'], errors='coerce')
    no_time = df_copy[df_copy['_dt'].isna()].drop(columns=['_dt'])
    has_time = df_copy[df_copy['_dt'].notna()]
    early = has_time[has_time['_dt'].dt.hour < before_hour].drop(columns=['_dt'])
    rest = has_time[has_time['_dt'].dt.hour >= before_hour].drop(columns=['_dt'])
    if config.TIME_OFF_WORK_COL in no_time.columns:
        no_time = no_time[no_time[config.TIME_OFF_WORK_COL].notna() & (no_time[config.TIME_OFF_WORK_COL].astype(str).str.strip() != '')]
    return pd.concat([rest, no_time], ignore_index=True), early





def _write_back_sheet(file_path, sheet_name, df):
    wb = openpyxl.load_workbook(file_path)
    ws = wb[sheet_name]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row)
    for r_idx, row in enumerate(df.itertuples(index=False), start=2):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val if not pd.isna(val) else '')
    wb.save(file_path)


def _update_cells(file_path, sheet_name, updates):
    wb = openpyxl.load_workbook(file_path)
    ws = wb[sheet_name]
    for row_idx, col_name, value in updates:
        for c_idx in range(1, ws.max_column + 1):
            if ws.cell(row=1, column=c_idx).value == col_name:
                ws.cell(row=row_idx, column=c_idx, value=value if not pd.isna(value) else '')
                break
    wb.save(file_path)


def _append_row(file_path, sheet_name, row_data):
    wb = openpyxl.load_workbook(file_path)
    ws = wb[sheet_name]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    new_row = []
    for h in headers:
        new_row.append(row_data.get(h, '') if h in row_data else '')
    ws.append(new_row)
    wb.save(file_path)





def read_excel_file(file_path):

    try:

        sheet_names = _find_target_sheets(file_path)

        if len(sheet_names) >= 2:

            print(f"发现 {len(sheet_names)} 个数据Sheet: {sheet_names}")

            df_day1 = pd.read_excel(file_path, sheet_name=sheet_names[0])
            df_day2 = pd.read_excel(file_path, sheet_name=sheet_names[1])

            id_col = '工号'
            wb = openpyxl.load_workbook(file_path)
            ws = wb[sheet_names[0]]
            merged_flag = ws.cell(row=1, column=ws.max_column).value
            wb.close()
            if merged_flag == '__MERGED__':
                print(f"Sheet1 已标记为合并数据，跳过合并")
                df_raw = df_day1
            else:
                print(f"Sheet1 ({sheet_names[0]}): {len(df_day1)} 行")
                print(f"Sheet2 ({sheet_names[1]}): {len(df_day2)} 行")

                df_day1, _ = _filter_by_checkin_time(df_day1, before_hour=9)
                df_day2_rest, df_day2_early = _filter_by_checkin_time(df_day2, before_hour=9)

                if config.TIME_OFF_WORK_COL in df_day2_rest.columns:
                    d2_no_checkin = df_day2_rest[df_day2_rest[config.TIME_OFF_WORK_COL].notna() & (df_day2_rest[config.TIME_OFF_WORK_COL].astype(str).str.strip() != '')]
                    no_time_mask = df_day2['打卡时间'].isna() if '打卡时间' in df_day2.columns else pd.Series([False] * len(df_day2))
                    d2_no_checkin = d2_no_checkin[no_time_mask.reindex(d2_no_checkin.index).fillna(False)]
                else:
                    d2_no_checkin = pd.DataFrame(columns=df_day2_rest.columns)
                df_day2_early = pd.concat([df_day2_early, d2_no_checkin], ignore_index=True)

                print(f"Sheet1 删除凌晨打卡后: {len(df_day1)} 行")
                print(f"Sheet2 凌晨打卡+无打卡: {len(df_day2_early)} 行")

                id_col = '工号'
                room_col = '房间号/工作点'
                d2_ids = df_day2_early[id_col].dropna().astype(str).tolist() if id_col in df_day2_early.columns else []
                d2_all_ids = df_day2[id_col].dropna().astype(str).tolist() if id_col in df_day2.columns else []

                if id_col in df_day1.columns:
                    if d2_ids:
                        df_day1 = df_day1[~df_day1[id_col].astype(str).isin(d2_ids)]
                        print(f"Sheet1 删除被Sheet2凌晨数据覆盖的行后: {len(df_day1)} 行")
                    if room_col in df_day1.columns and d2_all_ids:
                        is_handover = df_day1[room_col].astype(str).str.strip() == '接班'
                        not_in_d2 = ~df_day1[id_col].astype(str).isin(d2_all_ids)
                        has_off_work = df_day1[config.TIME_OFF_WORK_COL].notna() & (df_day1[config.TIME_OFF_WORK_COL].astype(str).str.strip() != '')
                        remove_mask = is_handover & not_in_d2 & ~has_off_work
                        if remove_mask.any():
                            print(f"Sheet1 删除接班但Sheet2无记录且无下班时间的行: {remove_mask.sum()} 行")
                            df_day1 = df_day1[~remove_mask]

                df_raw = pd.concat([df_day1, df_day2_early], ignore_index=True)
                print(f"合并后共 {len(df_raw)} 行数据")

                _write_back_sheet(file_path, sheet_names[0], df_raw)
                wb = openpyxl.load_workbook(file_path)
                ws = wb[sheet_names[0]]
                has_flag = False
                for c in range(1, ws.max_column + 1):
                    if ws.cell(row=1, column=c).value == '__MERGED__':
                        has_flag = True
                        break
                if not has_flag:
                    ws.cell(row=1, column=ws.max_column + 1, value='__MERGED__')
                    wb.save(file_path)
                wb.close()
                print(f"已将合并数据写回Sheet: {sheet_names[0]}")

        elif len(sheet_names) == 1:
            print(f"使用Sheet: {sheet_names[0]}")
            df_raw = pd.read_excel(file_path, sheet_name=sheet_names[0])
            df_raw, _ = _filter_by_checkin_time(df_raw, before_hour=9)
            print(f"删除凌晨打卡后: {len(df_raw)} 行")

        else:

            df_raw = pd.read_excel(file_path, sheet_name=0)

        print(f"成功读取文件，共 {len(df_raw)} 行数据")



        df_raw = df_raw[df_raw[config.TIME_OFF_WORK_COL] != '无']

        drop_cols = [c for c in df_raw.columns if str(c).startswith('__MERGED__')]
        if drop_cols:
            df_raw = df_raw.drop(columns=drop_cols)

        print(f"删除下班时间为'无'的数据后，剩余 {len(df_raw)} 行数据")



        invalid_time_data = []

        for idx, row in df_raw.iterrows():

            if pd.notna(row[config.TIME_OFF_WORK_COL]):

                time_str = str(row[config.TIME_OFF_WORK_COL])

                if not is_valid_time(time_str):

                    invalid_time_data.append({

                        '姓名': row['姓名'],

                        '工号': row['工号'],

                        '下班时间': time_str,

                        'Excel行号': idx + 2

                    })



        if invalid_time_data:

            msg = "\n发现不合法的时间格式！\n" + "=" * 50 + "\n以下记录的时间格式不正确：\n"

            msg += tabulate(pd.DataFrame(invalid_time_data), headers='keys', tablefmt='grid', showindex=False)

            raise ValueError(msg)



        duplicate_ids = df_raw[df_raw['工号'].notna()].groupby('工号').filter(lambda x: len(x) > 1)

        if not duplicate_ids.empty:

            msg = "\n发现重复的工号！\n" + "=" * 50 + "\n以下工号出现多次：\n"

            duplicate_df = duplicate_ids[['姓名', '工号', config.TIME_OFF_WORK_COL]].copy()

            duplicate_df['Excel行号'] = duplicate_ids.index + 2

            msg += tabulate(duplicate_df, headers='keys', tablefmt='grid', showindex=False)

            raise ValueError(msg)



        print("\nExcel文件的列名：")

        print("=" * 50)

        for col in df_raw.columns:

            print(f"- {col}")



        print("正在重命名列...")

        df = df_raw.copy()

        df.columns = df.columns.str.strip()

        df = df.rename(columns=config.EXCEL_COLUMN_RENAME)

        print("列名重命名完成")

        print(f"列名：{list(df.columns)}")



        if '工号' in df.columns:

            df['工号'] = df['工号'].apply(lambda x: str(int(x)) if pd.notna(x) and isinstance(x, float) else (str(int(x)) if pd.notna(x) and isinstance(x, (int,)) else x))



        for col in df.columns:

            df[col] = df[col].apply(convert_to_half_width)

            if '时间' in col:

                df[col] = df[col].apply(normalize_time_str)



        return df

    except Exception as e:

        print(f"读取文件时出错: {e}")

        return None





def _classify_by_keyword(location):

    if '接班' in location:

        return '接班'

    elif '值班' in location:

        return '值班'

    elif '休' in location:

        return '休息'

    elif any(kw in location for kw in config.OUTSIDE_KEYWORDS):

        return '外围'

    return None





def _classify_location(next_location):

    if not pd.notna(next_location) or next_location == 'nan':

        return '其他'



    next_location = str(next_location).strip()

    target_part = next_location



    if '+' in next_location:

        target_part = next_location.split('+')[-1].strip()



        keyword_cat = _classify_by_keyword(target_part)

        if keyword_cat:

            return keyword_cat



    match = re.match(r'^(\d+)(?:[ vV]\d+(?:\.\d+)?)?$', target_part)

    if not match:

        match = re.match(r'^(?:[\u4e00-\u9fffA-Za-z]+[+ ]?)?(\d+)(?:[ vV]\d+(?:\.\d+)?)?$', target_part)

    if match:

        room_num = int(match.group(1))

        for category, rooms in config.ROOM_RANGES.items():

            if room_num in rooms:

                return category



    keyword_cat = _classify_by_keyword(target_part)

    if keyword_cat:

        return keyword_cat



    return '其他'





def _sort_custom(data_list):

    handover = []

    normal = []

    for data in data_list:

        if pd.notna(data['当日']) and str(data['当日']).strip() == '接班':

            handover.append(data)

        else:

            normal.append(data)



    def work_hours_key(data):

        hours = data.get('上班时长(小时)')

        if isinstance(hours, (int, float)):

            return hours

        return float('inf')



    normal_sorted = sorted(normal, key=work_hours_key)



    def time_to_minutes(time_str):

        if time_str and ':' in str(time_str):

            try:

                time_part = str(time_str).split('(')[0].strip()

                h, m = map(int, time_part.split(':'))

                if h < 7:

                    h += 24

                return h * 60 + m

            except:

                return float('inf')

        return float('inf')



    for h in handover:

        h_time = time_to_minutes(h.get(config.TIME_OFF_WORK_COL, ''))

        inserted = False

        for i, n in enumerate(normal_sorted):

            n_time = time_to_minutes(n.get(config.TIME_OFF_WORK_COL, ''))

            if h_time < n_time:

                normal_sorted.insert(i, h)

                inserted = True

                break

        if not inserted:

            normal_sorted.append(h)

    return normal_sorted





def process_work_data(df, file_path, show_missing=True):

    if df.empty:

        return {}



    try:

        date_part = os.path.basename(file_path).split('_')[0]

        if '至' in date_part:

            date_str = date_part.split('至')[0]

        else:

            date_str = datetime.now().strftime('%Y-%m-%d')

    except:

        date_str = datetime.now().strftime('%Y-%m-%d')



    df = df.dropna(subset=['姓名'])



    result = {}

    for _, row in df.iterrows():

        employee_id = row.get('工号', '')

        if pd.isna(employee_id) or employee_id == '':

            continue



        employee_data = {

            '工号': employee_id,

            '姓名': row['姓名']

        }



        for col in df.columns:

            if col != '工号':

                employee_data[col] = row[col]



        employee_data['日期'] = date_str



        if pd.notna(employee_data[config.TIME_OFF_WORK_COL]):

            try:

                end_time_str = employee_data[config.TIME_OFF_WORK_COL]

                parsed_time = parse_time_string(end_time_str)



                if not parsed_time:

                    raise ValueError("无效的时间格式")



                end_time_str_main = parsed_time['main_time']



                if ':' in end_time_str_main:

                    try:

                        hours, minutes = end_time_str_main.split(':')

                        hours = int(hours)

                        minutes = int(minutes)

                        end_time = timedelta(hours=hours, minutes=minutes)

                    except ValueError:

                        employee_data['上班时长(小时)'] = '时间格式错误'

                        continue

                else:

                    employee_data['上班时长(小时)'] = '时间格式错误'

                    continue



                if parsed_time and parsed_time.get('start_time'):

                    start_time_str = parsed_time['start_time']

                    employee_data['上班时间'] = start_time_str

                    hours, minutes = start_time_str.split(':')

                    hours = int(hours)

                    minutes = int(minutes)

                    start_time = timedelta(hours=hours, minutes=minutes)

                else:

                    day_location = str(employee_data.get('当日', '')).strip()

                    day_parsed = parse_time_string(day_location) if pd.notna(day_location) and day_location != 'nan' else None



                    if day_parsed and day_parsed.get('start_time'):

                        start_time_str = day_parsed['start_time']

                        employee_data['上班时间'] = start_time_str

                        hours, minutes = start_time_str.split(':')

                        hours = int(hours)

                        minutes = int(minutes)

                        start_time = timedelta(hours=hours, minutes=minutes)

                    elif day_location == '接班':

                        h, m = map(int, config.HANDOVER_START_TIME.split(':'))

                        start_time = timedelta(hours=h, minutes=m)

                        employee_data['上班时间'] = config.HANDOVER_START_TIME

                    else:

                        h, m = map(int, config.DEFAULT_START_TIME.split(':'))

                        start_time = timedelta(hours=h, minutes=m)

                        employee_data['上班时间'] = config.DEFAULT_START_TIME



                time_diff = end_time - start_time

                if time_diff.total_seconds() < 0:

                    time_diff = time_diff + timedelta(hours=24)



                hours = round(time_diff.total_seconds() / 3600, 2)

                employee_data['上班时长(小时)'] = hours



            except Exception as e:

                employee_data['上班时长(小时)'] = '时间格式错误'

        else:

            employee_data['上班时长(小时)'] = '无下班时间'



        result[employee_id] = employee_data



    category_buckets = {cat: [] for cat in config.CATEGORY_ORDER}



    for employee_id, data in result.items():

        next_location = str(data.get('次日', '')).strip()

        cat = _classify_location(next_location)

        if cat in category_buckets:

            category_buckets[cat].append(data)

        else:

            category_buckets['其他'].append(data)



    for cat in category_buckets:

        category_buckets[cat] = _sort_custom(category_buckets[cat])



    room_ranges = config.ROOM_RANGES

    missing_rooms = {}

    for category, room_range in room_ranges.items():

        data_list = category_buckets.get(category, [])

        d2_rooms = []

        for data in data_list:

            d2_value = str(data['次日']).strip()

            if pd.notna(d2_value) and d2_value != 'nan':

                if re.match(r'^[\u4e00-\u9fffA-Za-z]+[+ ]?(\d+)$', d2_value):

                    num_match = re.search(r'(\d+)$', d2_value)

                    if num_match:

                        d2_rooms.append(int(num_match.group(1)))

                elif re.match(r'^\d+$', d2_value):

                    d2_rooms.append(int(d2_value))



        missing = [str(room) for room in room_range if room not in d2_rooms]

        if missing:

            missing_rooms[category] = ', '.join(missing)



    categories = [(cat, category_buckets[cat]) for cat in config.CATEGORY_ORDER]



    return {

        'date': date_str,

        'categories': categories,

        'missing_rooms': missing_rooms if show_missing else {}

    }





def generate_html_table(data):

    template = Template(config.HTML_TEMPLATE)



    for category, data_list in data['categories']:
        for record in data_list:
            if pd.notna(record.get('当日')) and str(record.get('当日')).strip() in ('接班', '值班'):
                record['css_class'] = ''
            elif pd.notna(record.get('次日')) and '夜休' in str(record.get('次日')):
                record['css_class'] = ''
            else:

                time_str = record.get(config.TIME_OFF_WORK_COL, '')

                work_hours = record.get('上班时长(小时)')

                time_category = get_time_category(time_str, work_hours)

                record['css_class'] = time_category if time_category != 'normal' else ''



    date_str = data['date']

    if '-' in date_str:

        date_obj = datetime.strptime(date_str, "%Y-%m-%d")

    else:

        date_obj = datetime.strptime(date_str, "%Y%m%d")

    title = f"{date_obj.year}-{date_obj.month:02d}-{date_obj.day:02d}(星期{config.WEEK_MAP[date_obj.weekday()]})"



    html_content = template.render(

        title=title,

        categories=data['categories'],

        missing_rooms=data['missing_rooms']

    )

    return html_content





def create_table_image_from_html(html_file_path, output_filename):

    if PLAYWRIGHT_AVAILABLE:

        try:

            print("正在使用playwright转换HTML为图片...")

            with sync_playwright() as p:

                if getattr(sys, 'frozen', False):

                    base_path = sys._MEIPASS

                    browser_path = os.path.join(base_path, 'ms-playwright', 'chromium-1148', 'chrome-win', 'chrome.exe')

                    browser = p.chromium.launch(headless=True, executable_path=browser_path) if os.path.exists(browser_path) else p.chromium.launch(headless=True)

                else:

                    browser = p.chromium.launch(headless=True)

                page = browser.new_page()

                page.goto(f"file:///{os.path.abspath(html_file_path).replace(os.sep, '/')}")

                page.wait_for_load_state('networkidle')

                page.screenshot(path=output_filename, full_page=True)

                browser.close()

            print(f"图片已保存到: {output_filename}")

            return output_filename

        except Exception as e:

            print(f"playwright转换失败: {e}")





def main(file_path=None, show_missing=True):

    if file_path is None:

        file_path = '2025-05-30至2025-05-31_打卡记录.xlsx'



    if not os.path.exists(file_path):

        print(f"文件 {file_path} 不存在")

        return



    print(f"正在读取文件：{file_path}")

    df = read_excel_file(file_path)



    if df is None or df.empty:

        print("没有找到有效的数据")

        return



    print("处理工作数据...")

    data = process_work_data(df, file_path, show_missing=show_missing)



    if not data:

        print("没有有效的工作数据")

        return



    print("生成HTML表格...")

    html_content = generate_html_table(data)



    html_filename = config.HTML_OUTPUT_FILENAME

    with open(html_filename, 'w', encoding='utf-8') as f:

        f.write(html_content)

    print(f"HTML文件已保存: {html_filename}")



    print("生成表格图片...")

    base_name = os.path.splitext(os.path.basename(file_path))[0]

    if '至' in base_name:

        date_str = base_name.split('至')[0].replace('-', '')

    else:

        date_str = datetime.now().strftime('%Y%m%d')



    image_filename = create_table_image_from_html(html_filename, f"{date_str}{config.IMAGE_SUFFIX}")

    print("处理完成！")



    return html_filename, image_filename





COL_ORIGINAL_D1 = '房间号/工作点'

COL_ORIGINAL_D2 = '第二天房间号/工作点（周五打卡填周一工作点）'





def read_excel_overview(file_path):
    sheet_names = _find_target_sheets(file_path)
    sheet = sheet_names[0] if sheet_names else 0
    if isinstance(sheet, int):
        xl = pd.ExcelFile(file_path)
        sheet = xl.sheet_names[sheet]
    df = pd.read_excel(file_path, sheet_name=sheet)

    cols = ['姓名', '工号', COL_ORIGINAL_D1, config.TIME_OFF_WORK_COL, COL_ORIGINAL_D2]
    existing_cols = [c for c in cols if c in df.columns]
    result = df[existing_cols].copy()
    result = result.where(result.notna(), '')
    result['_sheet'] = sheet
    result['_orig_idx'] = range(len(result))
    return result, existing_cols





def save_excel_overview(file_path, changes, overview_df, log=print):
    sheet_map = {}
    for idx_str, fields in changes.items():
        idx = int(idx_str)
        row = overview_df.iloc[idx]
        sn = row.get('_sheet', '')
        if sn not in sheet_map:
            sheet_map[sn] = []
        for col, val in fields.items():
            sheet_map[sn].append((idx, col, val))
        sheet_map[sn].append((idx, '打卡时间', None))

    all_sheet_names = _find_target_sheets(file_path)
    for sn, updates_raw in sheet_map.items():
        target_sheet = sn
        if not target_sheet:
            target_sheet = all_sheet_names[0] if all_sheet_names else 'Sheet1'
        wb = openpyxl.load_workbook(file_path)
        if target_sheet not in wb.sheetnames:
            wb.close()
            continue
        ws = wb[target_sheet]
        header_map = {}
        for c in range(1, ws.max_column + 1):
            h = ws.cell(row=1, column=c).value
            if h:
                header_map[h] = c
        source_df = pd.read_excel(file_path, sheet_name=target_sheet)
        for merged_idx, col_name, value in updates_raw:
            orig_idx = int(overview_df.iloc[merged_idx].get('_orig_idx', merged_idx))
            if orig_idx >= len(source_df):
                continue
            excel_row = orig_idx + 2
            c = header_map.get(col_name)
            if c:
                ws.cell(row=excel_row, column=c, value=value if not pd.isna(value) else '')
            c_time = header_map.get('打卡时间')
            if c_time and col_name != '打卡时间':
                ws.cell(row=excel_row, column=c_time, value=None)
        wb.save(file_path)
    log(f"已保存 {len(changes)} 条修改到: {file_path}")


def patch_excel_data(file_path, name, employee_id, off_work_time, d1, d2, log=print):
    sheet_names = _find_target_sheets(file_path)
    sheet = sheet_names[0] if sheet_names else 0
    if isinstance(sheet, int):
        xl = pd.ExcelFile(file_path)
        sheet = xl.sheet_names[sheet]
    df = pd.read_excel(file_path, sheet_name=sheet)

    mask = df['工号'].astype(str) == str(employee_id)
    if mask.any():
        idx = mask.idxmax()
        updates = [
            (idx + 2, config.TIME_OFF_WORK_COL, off_work_time),
            (idx + 2, COL_ORIGINAL_D1, d1),
            (idx + 2, COL_ORIGINAL_D2, d2),
            (idx + 2, '打卡时间', None),
        ]
        _update_cells(file_path, sheet, updates)
        log(f"已修改第 {idx + 2} 行: {name}({employee_id}) 下班时间={off_work_time}")
    else:
        new_row = {col: '' for col in df.columns}
        new_row['姓名'] = name
        new_row['工号'] = employee_id
        new_row[config.TIME_OFF_WORK_COL] = off_work_time
        new_row[COL_ORIGINAL_D1] = d1 if d1 else ''
        new_row[COL_ORIGINAL_D2] = d2 if d2 else ''
        if '打卡时间' in df.columns:
            new_row['打卡时间'] = None
        _append_row(file_path, sheet, new_row)
        log(f"已新增行: {name}({employee_id}) 下班时间={off_work_time}")
    log(f"文件已保存: {file_path}")





if __name__ == "__main__":

    file_path = None

    show_missing = True



    if len(sys.argv) > 1:

        file_path = sys.argv[1]

    if len(sys.argv) > 2:

        show_missing = sys.argv[2].lower() in ['true', '1', 'yes']



    main(file_path, show_missing)

